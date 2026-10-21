#!/usr/bin/env python3
"""Build the 10-page editable assignment 24 report from its Markdown source."""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "bao-cao-de24.md"
OUTPUT = ROOT / "bao-cao-de24.docx"
REPORT_MARKER = "<!-- DE24_REPORT_SOURCE: ten-page-v1 -->"
EXPECTED_PAGE_BREAKS = 9
EXPECTED_IMAGES = [
    "architecture.png",
    "screenshots/compose-startup.png",
    "screenshots/services-status.png",
    "screenshots/task-tracker-ui.png",
    "screenshots/task-create.png",
    "screenshots/task-filter.png",
    "screenshots/pgadmin-tasks.png",
    "screenshots/api-e2e-terminal.png",
    "screenshots/security-headers.png",
    "screenshots/prometheus-targets.png",
    "screenshots/prometheus-metric-query.png",
    "screenshots/grafana-dashboard.png",
    "screenshots/grafana-explore.png",
    "screenshots/loki-logql-verification.png",
    "screenshots/hardening-verification.png",
    "screenshots/networks-db-roles.png",
    "screenshots/unit-api-tests.png",
    "screenshots/persistence-restart.png",
]
IMAGE_WIDTHS = [4.36, 5.18, 4.98, 4.07, 4.34, 4.20, 5.79, 5.48, 6.26,
                2.94, 3.21, 3.38, 3.97, 4.83, 4.99, 4.96, 4.74, 4.32]
TABLE_WIDTHS = [(2.24, 4.45), (1.18, 5.51), (2.76, 3.94), (2.17, 4.53)]
FONT = "Times New Roman"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 italic: bool | None = None, color: str = "000000") -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, label: str, target: str) -> None:
    relationship = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for family in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{family}"), FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, size: float | None = None,
               base_bold: bool | None = None, base_italic: bool | None = None) -> None:
    """Render bold, italic, code and links without changing the report typeface."""
    token = re.compile(r"(\*\*.*?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*)|\[[^\]]+\]\(https?://[^)]+\))")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, bold=base_bold, italic=base_italic)
        item = match.group(0)
        if item.startswith("**"):
            run = paragraph.add_run(item[2:-2])
            set_run_font(run, size=size, bold=True, italic=base_italic)
        elif item.startswith("`"):
            run = paragraph.add_run(item[1:-1])
            set_run_font(run, size=size, bold=base_bold, italic=base_italic)
        elif item.startswith("["):
            label, target = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", item).groups()
            add_hyperlink(paragraph, label, target)
        else:
            run = paragraph.add_run(item[1:-1])
            set_run_font(run, size=size, bold=base_bold, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, bold=base_bold, italic=base_italic)


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = props.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        props.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=90, bottom=55, end=90) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9D9D9")


def parse_table(lines, doc, table_index: int) -> None:
    rows = []
    for line in lines:
        raw = line.strip().strip("|")
        cells = [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", raw)]
        if cells and all(re.fullmatch(r"[-: ]+", cell or "-") for cell in cells):
            continue
        rows.append(cells)
    if not rows or table_index >= len(TABLE_WIDTHS):
        raise ValueError("Unexpected table count or empty table in report source.")
    columns = max(map(len, rows))
    widths = TABLE_WIDTHS[table_index]
    if columns != len(widths):
        raise ValueError(f"Table {table_index + 1} has {columns} columns; expected {len(widths)}.")
    table = doc.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, cell in enumerate(cells):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            text = values[col_index] if col_index < len(values) else ""
            add_inline(p, text, size=10.5, base_bold=(row_index == 0))
            set_cell_shading(cell, "E5E7EB" if row_index == 0 else ("F5F6F8" if row_index % 2 == 0 else "FFFFFF"))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    set_table_borders(table)


def configure_styles(doc) -> None:
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)
    for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3 if name != "Title" else 0)
        style.paragraph_format.line_spacing = 1.0
    caption = doc.styles["Caption"]
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.line_spacing = 1.0


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, placeholder, end))
    set_run_font(run, size=9)


def configure_section(section) -> None:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.61)
    section.left_margin = Inches(0.87)
    section.right_margin = Inches(0.71)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.68), WD_TAB_ALIGNMENT.RIGHT)
    add_inline(header, "ĐẶNG QUỐC DOANH  •  DTC245200775", size=8.5)
    run = header.add_run("\t")
    set_run_font(run, size=8.5)
    add_inline(header, "TASK TRACKER  •  ĐỀ 24", size=8.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    add_page_field(footer)


def add_picture(doc, alt: str, relative_path: str, image_index: int) -> None:
    path = ROOT / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Missing tracked report image: {path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    inline = paragraph.add_run().add_picture(str(path), width=Inches(IMAGE_WIDTHS[image_index]))
    inline._inline.docPr.set("descr", alt)


def build() -> Path:
    text = SOURCE.read_text(encoding="utf-8-sig")
    if not text.startswith(REPORT_MARKER):
        raise ValueError("Report source marker is missing; refusing to build from an unsynchronized source.")
    lines = text.splitlines()
    page_breaks = sum(line.strip() == "<!-- PAGEBREAK -->" for line in lines)
    image_matches = [re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip()) for line in lines if line.strip().startswith("![")]
    image_paths = [match.group(2) for match in image_matches if match]
    if page_breaks != EXPECTED_PAGE_BREAKS:
        raise ValueError(f"Expected {EXPECTED_PAGE_BREAKS} manual page breaks; found {page_breaks}.")
    if image_paths != EXPECTED_IMAGES:
        raise ValueError("Report source must reference the 18 approved, tracked figures once and in order.")
    if "DTC245200775" not in text or "Đặng Quốc Doanh" not in text:
        raise ValueError("Report student identity is missing.")
    if not (ROOT / "architecture.png").is_file():
        raise FileNotFoundError("Tracked architecture image is missing.")

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    doc.core_properties.title = "Báo cáo Đề 24 — Hệ thống quản lý công việc Task Tracker"
    doc.core_properties.subject = "Triển khai và Quản trị Hệ thống Phần mềm"
    doc.core_properties.author = "Đặng Quốc Doanh"
    doc.core_properties.comments = "Đề 24; source Markdown và 18 hình minh chứng nằm trong repository."
    doc.core_properties.category = "Báo cáo học phần"

    cover = False
    image_index = 0
    table_index = 0
    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "<!-- COVER -->":
            cover = True
            i += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            cover = False
            i += 1
            continue
        if line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code, code_lines = True, []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for n, code in enumerate(code_lines):
                    if n:
                        p.add_run().add_break()
                    add_inline(p, code, size=10.5)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            parse_table(block, doc, table_index)
            table_index += 1
            continue
        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            add_picture(doc, image.group(1), image.group(2), image_index)
            image_index += 1
            i += 1
            continue
        if line.startswith("*") and line.endswith("*") and line.startswith("*Hình "):
            p = doc.add_paragraph(style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            add_inline(p, line[1:-1], size=9, base_italic=True)
            i += 1
            continue
        if line.startswith("# "):
            content = line[2:].strip()
            if cover:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(34)
                p.paragraph_format.space_after = Pt(8)
                add_inline(p, content, size=19, base_bold=True)
            else:
                p = doc.add_paragraph(style="Heading 1")
                add_inline(p, content, size=14, base_bold=True)
            i += 1
            continue
        if line.startswith("## "):
            content = line[3:].strip()
            if cover:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, content, size=17, base_bold=True)
            else:
                p = doc.add_paragraph(style="Heading 2")
                add_inline(p, content, size=12, base_bold=True)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cover else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1 if cover else 3)
            add_inline(p, line[4:].strip(), size=15 if cover else 11.5, base_bold=True)
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cover else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        if cover and line.startswith("TRƯỜNG "):
            p.paragraph_format.space_before = Pt(12)
        add_inline(p, line, size=11.5)
        i += 1

    if in_code:
        raise ValueError("Unclosed fenced code block in report source.")
    if image_index != len(EXPECTED_IMAGES) or table_index != len(TABLE_WIDTHS):
        raise ValueError(f"Builder consumed {image_index} figures and {table_index} tables unexpectedly.")
    doc.save(OUTPUT)
    validate_output(OUTPUT)
    print(f"Wrote {OUTPUT} with 18 tracked images and 9 explicit page breaks (10 planned pages).")
    return OUTPUT


def validate_output(path: Path) -> None:
    built = Document(path)
    if len(built.inline_shapes) != 18:
        raise ValueError(f"Built DOCX contains {len(built.inline_shapes)} images; expected 18.")
    runs = [r for p in built.paragraphs for r in p.runs if r.text.strip()]
    runs.extend(r for t in built.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip())
    wrong_fonts = [r.text for r in runs if r.font.name != FONT]
    if wrong_fonts:
        raise ValueError(f"Found editable text not explicitly set in {FONT}: {wrong_fonts[:3]}")
    content = "\n".join(p.text for p in built.paragraphs)
    if "Đặng Quốc Doanh" not in content:
        raise ValueError("Built DOCX is missing the student identity.")
    if "12 ca Unit/API" not in content or "13 ca" not in content:
        raise ValueError("Current 13-test count and historical 12-test image note must both be present.")
    if "Kết quả runtime Windows mới nhất do người dùng cung cấp" not in content:
        raise ValueError("User-provided current Windows runtime results are missing from the report.")


if __name__ == "__main__":
    build()
